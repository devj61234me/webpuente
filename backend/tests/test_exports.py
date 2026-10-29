import io
import unittest
import json

from openpyxl import load_workbook

from services import exporters

from .base import ApiTest, error_of


class ExportTests(ApiTest):
    def setUp(self):
        super().setUp()
        self.project = self.analyzed_project()
        self.pid = self.project["id"]

    def approve(self, suffix, **body):
        row = self.mapping_for(self.pid, suffix)
        r = self.api.post(f"/api/projects/{self.pid}/mappings/{row['id']}/approve/", {"confirm_low_confidence": True, **body}, format="json")
        self.assertEqual(r.status_code, 200, r.content)
        return row

    def export(self, fmt, expect=201, **body):
        r = self.api.post(f"/api/projects/{self.pid}/exports/", {"format": fmt, **body}, format="json")
        self.assertEqual(r.status_code, expect, r.content)
        return r.json()

    def download(self, job):
        r = self.api.get(f"/api/projects/{self.pid}/exports/{job['id']}/download/")
        self.assertEqual(r.status_code, 200)
        return r

    def test_nothing_approved_means_nothing_exported(self):
        self.assertEqual(error_of(self.api.post(f"/api/projects/{self.pid}/exports/", {"format": "csv"}, format="json"))["code"], "nothing_to_export")

    def test_invalid_format(self):
        self.assertEqual(error_of(self.api.post(f"/api/projects/{self.pid}/exports/", {"format": "pdf"}, format="json"))["code"], "invalid_format")

    def test_csv_contains_only_approved_rows(self):
        self.approve("/about")
        job = self.export("csv")
        self.assertEqual(job["row_count"], 1)
        response = self.download(job)
        text = response.content.decode("utf-8-sig")
        self.assertIn("https://new.example.com/about", text)
        self.assertNotIn("trail-guide", text)
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertEqual(response["X-Content-Type-Options"], "nosniff")

    def test_pending_rows_need_explicit_confirmation(self):
        self.approve("/about")
        r = self.api.post(f"/api/projects/{self.pid}/exports/", {"format": "csv", "include_pending": True}, format="json")
        self.assertEqual(error_of(r)["code"], "pending_confirmation_required")
        job = self.export("csv", include_pending=True, confirm_pending=True)
        self.assertGreater(job["row_count"], 1)
        self.assertTrue(any("NOT been approved" in w for w in job["warnings"]))

    def test_xlsx_opens_and_has_a_summary_sheet(self):
        self.approve("/about")
        wb = load_workbook(io.BytesIO(self.download(self.export("xlsx")).content))
        self.assertGreaterEqual(len(wb.sheetnames), 2)

    def test_json_is_valid(self):
        self.approve("/about")
        data = json.loads(self.download(self.export("json")).content)
        self.assertIn("redirects", data) if "redirects" in data else self.assertTrue(data)

    def test_formula_injection_is_neutralised(self):
        self.assertTrue(exporters.safe_cell("=HYPERLINK(\"http://evil\")").startswith("'"))
        for bad in ("+1+1", "-2", "@SUM(A1)", "\t=1", "\r=1"):
            self.assertTrue(exporters.safe_cell(bad).startswith("'"), bad)
        self.assertEqual(exporters.safe_cell("normal text"), "normal text")
        row = self.approve("/about")
        self.api.patch(f"/api/projects/{self.pid}/mappings/{row['id']}/", {"note": "=cmd|' /C calc'!A0"}, format="json")
        text = self.download(self.export("csv")).content.decode("utf-8-sig")
        self.assertNotIn(",=cmd", text)
        self.assertIn("'=cmd", text)

    def test_htaccess_and_nginx_are_generated_and_pass_the_format_check(self):
        self.approve("/about")
        self.approve("/products/running-shoes", confirm_many_to_one=True)
        ht = self.download(self.export("htaccess")).content.decode()
        self.assertIn("RewriteRule ^about/?$", ht); self.assertIn("[R=301,L,NE]", ht)
        self.assertIn("https://new.example.com/about", ht)
        ng = self.download(self.export("nginx")).content.decode()
        self.assertIn("map $request_uri $webpuente_r301", ng)
        self.assertEqual(exporters.validate_syntax("nginx", ng.encode()), [])
        self.assertEqual(exporters.validate_syntax("htaccess", ht.encode()), [])

    def test_dangerous_characters_in_destinations_are_encoded(self):
        self.assertNotIn('"', exporters.safe_destination('https://new.example.com/a"b'))
        self.assertNotIn("$", exporters.safe_destination("https://new.example.com/a$b"))
        self.assertNotIn(" ", exporters.safe_destination("https://new.example.com/a b"))
        self.assertNotIn("\\", exporters.safe_destination("https://new.example.com/a\\b"))

    def test_status_override_and_validation(self):
        self.approve("/about")
        ht = self.download(self.export("htaccess", redirect_status=302)).content.decode()
        self.assertIn("[R=302,L,NE]", ht)
        r = self.api.post(f"/api/projects/{self.pid}/exports/", {"format": "csv", "redirect_status": 200}, format="json")
        self.assertEqual(error_of(r)["code"], "invalid_status")

    def test_config_exports_skip_problem_rows_and_say_so(self):
        row = self.mapping_for(self.pid, "/legacy/zzz-unrelated")
        self.api.post(f"/api/projects/{self.pid}/mappings/{row['id']}/assign/", {"url": "https://new.example.com/ok"}, format="json")
        self.approve("/about")
        job = self.export("nginx")
        self.assertGreaterEqual(job["row_count"], 1)
        self.assertTrue(any("server" in w for w in job["warnings"]))

    def test_exports_are_listed_and_private_to_the_owner(self):
        self.approve("/about")
        job = self.export("csv")
        self.assertEqual(self.api.get(f"/api/projects/{self.pid}/exports/").json()["count"], 1)
        other = self.client_for(self.make_user("other@example.com"))
        self.assertEqual(other.get(f"/api/projects/{self.pid}/exports/{job['id']}/download/").status_code, 404)

    def test_report_includes_all_statuses_even_without_approvals(self):
        job = self.export("report")
        self.assertEqual(job["row_count"], 6)


# ---------------------------------------------------------------------------------------------------------
# Generated server rules, checked without the database. If a real nginx is installed they are also loaded by it.
import shutil
import socket
import subprocess
import tempfile
import time
import urllib.request
from pathlib import Path

from django.test import SimpleTestCase

from services.urlnorm import NormRules, normalize_url

RULES = NormRules.from_dict({"ignore_case": False, "strip_trailing_slash": True, "query_mode": "keep", "drop_params": [], "ignore_www": False})


def export_rows(cases):
    rows = []
    for old, new, code in cases:
        o, n = normalize_url(old, rules=RULES), normalize_url(new, rules=RULES)
        assert o.ok and n.ok, (old, new, o.error, n.error)
        rows.append({"old_path_key": o.path_key, "new_normalized": n.full, "redirect_status": code})
    return rows


NGINX_CASES = [
    ("https://old.example.com/about", "https://new.example.com/about-us", 301),
    ("https://old.example.com/a%20b/c", "https://new.example.com/x%20y", 301),
    ("https://old.example.com/price$list", "https://new.example.com/p$1?x=$y", 301),
    ("https://old.example.com/q%22uote", "https://new.example.com/say%22hi%22", 302),
    ("https://old.example.com/caf%C3%A9", "https://new.example.com/café", 301),
    ("https://old.example.com/search?page=2&b=1", "https://new.example.com/find", 301),
]


class NginxRuleTests(SimpleTestCase):
    def test_dollar_in_the_old_path_is_matched_raw_and_encoded_but_never_in_the_destination(self):
        text = exporters.build_nginx(export_rows(NGINX_CASES), "t", True).decode()
        self.assertIn('"/price$list" "https://new.example.com/p%241?x=%24y";', text)
        self.assertIn('"/price%24list" "https://new.example.com/p%241?x=%24y";', text)
        for line in text.splitlines():
            if line.strip().startswith('"'):
                value = line.split('" "', 1)[1]
                self.assertNotIn("$", value, line)  # nginx would expand a $ inside a value
        self.assertEqual(exporters.validate_syntax("nginx", text.encode()), [])

    @unittest.skipUnless(shutil.which("nginx"), "nginx is not installed on this machine")
    def test_real_nginx_accepts_the_file_and_redirects_correctly(self):
        with socket.socket() as probe:
            probe.bind(("127.0.0.1", 0))
            port = probe.getsockname()[1]
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            (tmp / "map.conf").write_bytes(exporters.build_nginx(export_rows(NGINX_CASES), "t", True))
            (tmp / "nginx.conf").write_text(
                f"pid {tmp}/nginx.pid;\nerror_log {tmp}/error.log;\nevents {{}}\nhttp {{\n  access_log off;\n"
                f"  client_body_temp_path {tmp}/b; proxy_temp_path {tmp}/p; fastcgi_temp_path {tmp}/f; uwsgi_temp_path {tmp}/u; scgi_temp_path {tmp}/s;\n"
                f"  include {tmp}/map.conf;\n  server {{\n    listen 127.0.0.1:{port};\n"
                "    if ($webpuente_r301) { return 301 $webpuente_r301; }\n    if ($webpuente_r302) { return 302 $webpuente_r302; }\n"
                '    location / { return 200 "none"; }\n  }\n}\n'
            )
            test = subprocess.run(["nginx", "-t", "-c", str(tmp / "nginx.conf")], capture_output=True, text=True)
            self.assertEqual(test.returncode, 0, test.stderr)
            subprocess.run(["nginx", "-c", str(tmp / "nginx.conf")], check=True)
            try:
                for _ in range(50):
                    try:
                        socket.create_connection(("127.0.0.1", port), timeout=0.2).close()
                        break
                    except OSError:
                        time.sleep(0.1)

                class NoFollow(urllib.request.HTTPRedirectHandler):
                    def redirect_request(self, *args, **kwargs):
                        return None

                opener = urllib.request.build_opener(NoFollow)

                def get(path):
                    try:
                        with opener.open(f"http://127.0.0.1:{port}{path}") as r:
                            return r.status, ""
                    except urllib.error.HTTPError as e:
                        return e.code, e.headers.get("Location", "")

                self.assertEqual(get("/about"), (301, "https://new.example.com/about-us"))
                self.assertEqual(get("/about/"), (301, "https://new.example.com/about-us"))
                self.assertEqual(get("/a%20b/c"), (301, "https://new.example.com/x%20y"))
                self.assertEqual(get("/price$list"), (301, "https://new.example.com/p%241?x=%24y"))
                self.assertEqual(get("/price%24list"), (301, "https://new.example.com/p%241?x=%24y"))
                self.assertEqual(get("/q%22uote"), (302, "https://new.example.com/say%22hi%22"))
                self.assertEqual(get("/caf%C3%A9"), (301, "https://new.example.com/caf%C3%A9"))
                self.assertEqual(get("/search?b=1&page=2"), (301, "https://new.example.com/find"))
                self.assertEqual(get("/not-listed"), (200, ""))
            finally:
                subprocess.run(["nginx", "-c", str(tmp / "nginx.conf"), "-s", "stop"], capture_output=True)
                time.sleep(0.3)


APACHE_CASES = [
    ("https://old.example.com/about", "https://new.example.com/about-us", 301),
    ("https://old.example.com/a%20b/c", "https://new.example.com/x%20y", 301),
    ("https://old.example.com/price$list", "https://new.example.com/p$1?x=$y", 301),
    ("https://old.example.com/q%22uote", "https://new.example.com/say%22hi%22", 302),
    ("https://old.example.com/caf%C3%A9", "https://new.example.com/café", 301),
    ("https://old.example.com/a.b+c(d)", "https://new.example.com/regex-chars", 301),
    ("https://old.example.com/", "https://new.example.com/home", 301),
    ("https://old.example.com/search", "https://new.example.com/search-plain", 301),
    ("https://old.example.com/search?page=2&b=1", "https://new.example.com/find%20it", 301),
]
APACHE_MODULES = Path("/usr/lib/apache2/modules")


class ApacheRuleTests(SimpleTestCase):
    def test_destinations_cannot_be_double_encoded_or_back_referenced(self):
        text = exporters.build_htaccess(export_rows(APACHE_CASES), "t").decode()
        self.assertIn('"https://new.example.com/x\\%20y"', text)         # every % is escaped for mod_rewrite
        self.assertIn('"https://new.example.com/p\\%241?x=\\%24y"', text)  # and a $ never survives into the destination
        self.assertNotRegex(text, r'RewriteRule \S+ "[^"]*[^\\]%[0-9A-Fa-f]{2}')  # no unescaped %XX in any substitution
        self.assertEqual(exporters.validate_syntax("htaccess", text.encode()), [])

    def test_query_rules_come_before_plain_rules_for_the_same_path(self):
        text = exporters.build_htaccess(export_rows(APACHE_CASES), "t").decode()
        self.assertLess(text.index("find\\%20it"), text.index("search-plain"))

    def test_the_validator_accepts_escaped_quotes_but_rejects_garbage(self):
        good = b'<IfModule mod_rewrite.c>\nRewriteEngine On\nRewriteRule ^q\\"uote/?$ "https://x/y" [R=301,L,NE]\n</IfModule>\n'
        self.assertEqual(exporters.validate_syntax("htaccess", good), [])
        for bad in (b"RedirectMatch 301 ^/a https://x/y\n", b"RewriteRule evil [R=301]\n", b"Header set X evil\n"):
            self.assertTrue(exporters.validate_syntax("htaccess", bad), bad)

    @unittest.skipUnless(shutil.which("apache2") and (APACHE_MODULES / "mod_rewrite.so").exists(), "Apache (apache2 with mod_rewrite) is not installed here")
    def test_real_apache_redirects_exactly_as_intended(self):
        import os
        import pwd

        with socket.socket() as probe:
            probe.bind(("127.0.0.1", 0))
            port = probe.getsockname()[1]
        root = Path(tempfile.mkdtemp(prefix="wp-apache-"))
        try:
            os.chmod(root, 0o777)
            (root / "www").mkdir()
            os.chmod(root / "www", 0o755)
            (root / "www" / "index.html").write_text("ok")
            (root / "www" / ".htaccess").write_bytes(exporters.build_htaccess(export_rows(APACHE_CASES), "t"))
            for f in (root / "www").iterdir():
                os.chmod(f, 0o644)
            user = "User www-data\nGroup www-data\n" if os.geteuid() == 0 else ""
            if user:
                try:
                    pwd.getpwnam("www-data")
                except KeyError:
                    self.skipTest("no www-data user to run Apache as")
            mods = "".join(f"LoadModule {n}_module {APACHE_MODULES}/mod_{n}.so\n" for n in ("mpm_prefork", "authz_core", "rewrite", "dir", "mime"))
            (root / "apache.conf").write_text(
                f'ServerRoot "/etc/apache2"\nPidFile {root}/apache.pid\nServerName localhost\n{user}Listen 127.0.0.1:{port}\nErrorLog {root}/error.log\n{mods}'
                f"TypesConfig /etc/mime.types\nDocumentRoot {root}/www\n<Directory {root}/www>\n  AllowOverride All\n  Require all granted\n</Directory>\n"
            )
            check = subprocess.run(["apache2", "-t", "-f", str(root / "apache.conf")], capture_output=True, text=True)
            self.assertEqual(check.returncode, 0, check.stderr)
            subprocess.run(["apache2", "-f", str(root / "apache.conf")], check=True, capture_output=True)
            try:
                for _ in range(50):
                    try:
                        socket.create_connection(("127.0.0.1", port), timeout=0.2).close()
                        break
                    except OSError:
                        time.sleep(0.1)

                class NoFollow(urllib.request.HTTPRedirectHandler):
                    def redirect_request(self, *args, **kwargs):
                        return None

                opener = urllib.request.build_opener(NoFollow)

                def get(path):
                    try:
                        with opener.open(f"http://127.0.0.1:{port}{path}") as r:
                            return r.status, ""
                    except urllib.error.HTTPError as e:
                        return e.code, e.headers.get("Location", "")

                self.assertEqual(get("/about"), (301, "https://new.example.com/about-us"))
                self.assertEqual(get("/about/"), (301, "https://new.example.com/about-us"))
                self.assertEqual(get("/a%20b/c"), (301, "https://new.example.com/x%20y"))               # was /x%2520y
                self.assertEqual(get("/price$list"), (301, "https://new.example.com/p%241?x=%24y"))     # was /p%25241
                self.assertEqual(get("/price%24list"), (301, "https://new.example.com/p%241?x=%24y"))
                self.assertEqual(get("/q%22uote"), (302, "https://new.example.com/say%22hi%22"))
                self.assertEqual(get("/caf%C3%A9"), (301, "https://new.example.com/caf%C3%A9"))
                self.assertEqual(get("/a.b+c(d)"), (301, "https://new.example.com/regex-chars"))
                self.assertEqual(get("/aXb+c(d)")[0], 404)  # regex characters are escaped, so "." is not a wildcard
                self.assertEqual(get("/"), (301, "https://new.example.com/home"))
                self.assertEqual(get("/search"), (301, "https://new.example.com/search-plain"))
                self.assertEqual(get("/search?b=1&page=2"), (301, "https://new.example.com/find%20it"))  # query rule wins over plain
                self.assertEqual(get("/not-listed")[0], 404)
            finally:
                subprocess.run(["apache2", "-f", str(root / "apache.conf"), "-k", "stop"], capture_output=True)
                time.sleep(0.5)
        finally:
            shutil.rmtree(root, ignore_errors=True)
