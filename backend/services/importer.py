"""
Reading URL lists from CSV, XLSX or pasted text.

Two steps, so the user stays in control:
  1. preview_*  -> detect the header row, suggest which column is which, show a few sample rows
  2. parse_*    -> read every row with the (possibly corrected) column map; invalid rows are reported, never dropped silently
"""
import csv
import hashlib
import io
import re
from dataclasses import dataclass, field

from django.conf import settings

from core.exceptions import ServiceError

FIELDS = ("url", "title", "meta_description", "h1", "canonical_url", "status_code")

# Header names seen in common crawler exports (Screaming Frog, Sitebulb, spreadsheets people make by hand).
HEADER_HINTS = {
    "url": ["url", "address", "old url", "new url", "old_url", "new_url", "page url", "link", "source", "destination", "loc"],
    "title": ["title", "title 1", "page title", "meta title"],
    "meta_description": ["meta description", "meta description 1", "description", "meta"],
    "h1": ["h1", "h1-1", "h1 1", "heading", "heading 1"],
    "canonical_url": ["canonical", "canonical link element 1", "canonical url", "canonical link"],
    "status_code": ["status code", "status", "http status", "response code"],
}
URL_LIKE = re.compile(r"^(https?://|/)\S+$", re.IGNORECASE)


@dataclass
class ParsedRow:
    row_number: int
    values: dict  # keys from FIELDS (missing ones are absent)


@dataclass
class ParsedTable:
    headers: list
    has_header: bool
    rows: list = field(default_factory=list)       # list of lists of strings (data rows only)
    first_row_number: int = 1                      # row number of rows[0] in the file


def _clean(cell):
    return "" if cell is None else str(cell).strip()


def _decode(data):
    """UTF-8 (with or without BOM) first; fall back to Windows-1252 because Excel often saves CSV that way."""
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ServiceError("This file could not be read as text. Save it as CSV (UTF-8) and try again.", code="file_unreadable")


def check_size(data, filename=""):
    limit = settings.MAX_UPLOAD_MB * 1024 * 1024
    if len(data) > limit:
        raise ServiceError(f"The file is larger than {settings.MAX_UPLOAD_MB} MB.", code="file_too_large", status_code=413)
    if not data:
        raise ServiceError("The file is empty.", code="file_empty")


def file_kind(filename, data):
    """Decide the format from the extension AND the content (a renamed .exe must not pass)."""
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        if not data.startswith(b"PK"):  # an .xlsx is a zip file
            raise ServiceError("This file is not a valid .xlsx workbook.", code="file_invalid")
        return "xlsx"
    if name.endswith((".csv", ".txt", ".tsv")):
        if b"\x00" in data[:4096]:
            raise ServiceError("This file does not look like a text/CSV file.", code="file_invalid")
        return "csv"
    raise ServiceError("Only .csv and .xlsx files are supported.", code="file_type_not_allowed", status_code=400)


def _read_csv(data):
    text = _decode(data)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = []
    for row in reader:
        rows.append([_clean(c) for c in row])
        if len(rows) > settings.MAX_IMPORT_ROWS + 1:
            raise ServiceError(f"The file has more than {settings.MAX_IMPORT_ROWS} rows.", code="too_many_rows", status_code=413)
    return rows


def _read_xlsx(data):
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)  # values only, formulas are not run
    except Exception:  # noqa: BLE001 - any parse failure is "bad file" to the user
        raise ServiceError("This .xlsx file could not be opened.", code="file_invalid")
    sheet = workbook.worksheets[0]
    rows = []
    for row in sheet.iter_rows(values_only=True):
        rows.append([_clean(c) for c in row])
        if len(rows) > settings.MAX_IMPORT_ROWS + 1:
            workbook.close()
            raise ServiceError(f"The sheet has more than {settings.MAX_IMPORT_ROWS} rows.", code="too_many_rows", status_code=413)
    workbook.close()
    return rows


def _rows_from_text(text):
    lines = [line.strip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    if len([line for line in lines if line]) > settings.MAX_IMPORT_ROWS:
        raise ServiceError(f"You pasted more than {settings.MAX_IMPORT_ROWS} lines.", code="too_many_rows", status_code=413)
    return [[line] for line in lines if line]


def _looks_like_header(first_row):
    """A header row has no URL-looking cell. A list that starts straight with a URL has no header."""
    return not any(URL_LIKE.match(cell) for cell in first_row if cell)


def build_table(rows):
    rows = [r for r in rows if any(c for c in r)] if rows else []
    if not rows:
        raise ServiceError("No rows were found.", code="file_empty")
    if _looks_like_header(rows[0]):
        return ParsedTable(headers=rows[0], has_header=True, rows=rows[1:], first_row_number=2)
    width = max(len(r) for r in rows)
    return ParsedTable(headers=[f"Column {i + 1}" for i in range(width)], has_header=False, rows=rows, first_row_number=1)


def suggest_mapping(headers, rows):
    """Guess which column holds which field. Returns {field: column_index}. The user can change it."""
    mapping = {}
    lowered = [h.strip().lower() for h in headers]
    for field_name, hints in HEADER_HINTS.items():
        for hint in hints:
            if hint in lowered:
                mapping[field_name] = lowered.index(hint)
                break
    if "url" not in mapping:  # fall back to the first column whose cells look like URLs
        for index in range(len(headers)):
            cells = [r[index] for r in rows[:20] if index < len(r) and r[index]]
            if cells and sum(bool(URL_LIKE.match(c)) for c in cells) >= max(1, len(cells) // 2):
                mapping["url"] = index
                break
    return mapping


def table_from_upload(data, filename):
    check_size(data, filename)
    kind = file_kind(filename, data)
    rows = _read_xlsx(data) if kind == "xlsx" else _read_csv(data)
    return build_table(rows), hashlib.sha256(data).hexdigest()


def table_from_text(text):
    """Pasted text is always ONE URL PER LINE (no header row), so a wrong line is reported as invalid, not mistaken for a header."""
    if len(text) > settings.MAX_UPLOAD_MB * 1024 * 1024:
        raise ServiceError(f"The pasted text is larger than {settings.MAX_UPLOAD_MB} MB.", code="file_too_large", status_code=413)
    rows = _rows_from_text(text)
    if not rows:
        raise ServiceError("No rows were found.", code="file_empty")
    return ParsedTable(headers=["URL"], has_header=False, rows=rows, first_row_number=1), hashlib.sha256(text.encode("utf-8")).hexdigest()


def preview(table, sample_size=8):
    return {
        "headers": table.headers,
        "has_header": table.has_header,
        "total_rows": len(table.rows),
        "sample": table.rows[:sample_size],
        "suggested_mapping": suggest_mapping(table.headers, table.rows),
    }


def rows_with_mapping(table, mapping):
    """Yield ParsedRow objects using the column map {field: column_index}."""
    if "url" not in mapping or mapping["url"] is None:
        raise ServiceError("Choose which column contains the URL.", code="url_column_required", details={"column_map": ["Choose the URL column."]})
    width = len(table.headers)
    for field_name, index in mapping.items():
        if field_name not in FIELDS:
            raise ServiceError(f"Unknown field '{field_name}'.", code="invalid_column_map")
        if index is None:
            continue
        if not isinstance(index, int) or index < 0 or index >= max(width, 1):
            raise ServiceError("The column choice does not match the file.", code="invalid_column_map")
    for offset, row in enumerate(table.rows):
        values = {}
        for field_name, index in mapping.items():
            if index is not None and index < len(row) and row[index] != "":
                values[field_name] = row[index]
        yield ParsedRow(table.first_row_number + offset, values)
