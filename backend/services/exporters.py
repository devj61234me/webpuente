"""
Building the export files.

Security points:
- Spreadsheet cells that start with = + - @ (or tab/CR) are prefixed with an apostrophe, so Excel never runs them.
- Text going into .htaccess / nginx files is escaped and any character that could break out of a rule is encoded.
- Config files only contain approved, valid redirects. Skipped rows are reported, never silently dropped.
"""
import csv
import io
import json
import re
from datetime import datetime, timezone
from fnmatch import fnmatchcase
from urllib.parse import quote, unquote

SHEET_COLUMNS = [
    ("old_url", "Old URL"), ("old_normalized", "Old URL (normalised)"), ("new_url", "New URL"),
    ("new_normalized", "New URL (normalised)"), ("redirect_status", "Redirect status"), ("confidence", "Confidence"),
    ("band", "Confidence level"), ("match_type", "Match type"), ("reasons", "Reasons"),
    ("review_status", "Review status"), ("origin", "Source"), ("note", "Note"),
]
FORMATS = ("csv", "xlsx", "json", "htaccess", "nginx", "report")
CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "json": "application/json",
    "htaccess": "text/plain; charset=utf-8",
    "nginx": "text/plain; charset=utf-8",
    "report": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
EXTENSIONS = {"csv": "csv", "xlsx": "xlsx", "json": "json", "htaccess": "htaccess.txt", "nginx": "nginx.conf", "report": "xlsx"}


def safe_cell(value):
    """Stop spreadsheet formula injection."""
    if value is None:
        return ""
    text = str(value)
    if text and text[0] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + text
    return text


# ---------------------------------------------------------------- spreadsheet-like formats
def _sheet_rows(rows):
    for r in rows:
        yield [safe_cell(r.get(key)) if key != "reasons" else safe_cell("; ".join(r.get("reasons") or [])) for key, _ in SHEET_COLUMNS]


def build_csv(rows):
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow([label for _, label in SHEET_COLUMNS])
    writer.writerows(_sheet_rows(rows))
    return ("﻿" + buffer.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 correctly


def build_xlsx(rows, summary=None, risks=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    sheet = book.active
    sheet.title = "Mappings"
    sheet.append([label for _, label in SHEET_COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for line in _sheet_rows(rows):
        sheet.append(line)
    sheet.freeze_panes = "A2"
    if summary:
        info = book.create_sheet("Summary")
        for key, value in summary.items():
            info.append([safe_cell(key), safe_cell(value)])
        info.column_dimensions["A"].width = 32
        info.column_dimensions["B"].width = 60
    if risks is not None:
        sheet_risks = book.create_sheet("Risks")
        sheet_risks.append(["Severity", "Type", "Source", "Old URL", "New URL", "Message"])
        for cell in sheet_risks[1]:
            cell.font = Font(bold=True)
        for risk in risks:
            sheet_risks.append([safe_cell(risk.get(k)) for k in ("severity", "kind", "source", "old_url", "new_url", "message")])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def build_json(rows, meta):
    payload = {"generated_at": datetime.now(timezone.utc).isoformat(), **meta, "redirects": rows}
    return json.dumps(payload, indent=2, ensure_ascii=False).encode("utf-8")


# ---------------------------------------------------------------- server configuration files
def safe_destination(url):
    """Make a destination URL safe to place inside a server rule. $ " \\ and spaces are percent-encoded
    (a literal $ would be read as a regex back-reference; a quote would end the string)."""
    return re.sub(r'[$"\\\s<>{}|^`]', lambda m: quote(m.group(0), safe=""), url)


def _path_of(path_key):
    return path_key.split("?")[0]


def _apache_regex(decoded_path):
    """Escape every character that is special in a regular expression."""
    return re.sub(r"([^A-Za-z0-9/_-])", r"\\\1", decoded_path)


def _apache_comment(text):
    return re.sub(r"[\r\n]", " ", text)[:120]


def _apache_target(destination):
    """The substitution text of a RewriteRule. Verified against Apache 2.4.58:
      * every % must be written \\% - otherwise mod_rewrite reads %2 (as in %20) as a back-reference and deletes it;
      * the NE flag stops Apache from encoding the destination a second time (%20 would become %2520)."""
    return safe_destination(destination).replace("%", "\\%")


def build_htaccess(rows, project_name):
    """Apache rules (mod_rewrite, in the site's top folder). URLs that have a query string get their own, more specific
    rules, written FIRST so that a plain rule for the same path cannot shadow them."""
    lines = [
        "# WebPuente redirect rules - REVIEW BEFORE USE (see the checklist in the project docs).",
        f"# Project: {_apache_comment(project_name)}",
        f"# Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}",
        "# Needs Apache mod_rewrite. Put these lines ABOVE other rewrite rules, in the .htaccess of the site's top folder.",
        "",
        "<IfModule mod_rewrite.c>",
        "RewriteEngine On",
    ]
    plain = [r for r in rows if "?" not in r["old_path_key"]]
    with_query = [r for r in rows if "?" in r["old_path_key"]]

    def pattern_for(path):
        # In .htaccess the rewrite pattern is matched against the decoded path WITHOUT its leading slash
        return "^$" if path.strip("/") == "" else f"^{_apache_regex(path.strip('/'))}/?$"

    if with_query:
        lines.append("# URLs with a query string. The query must match exactly as written (parameters sorted A-Z).")
        for r in with_query:
            path, _, query = r["old_path_key"].partition("?")
            lines.append(f'RewriteCond %{{QUERY_STRING}} "^{_apache_regex(query)}$"')  # QUERY_STRING is the raw, undecoded text
            lines.append(f'RewriteRule {pattern_for(unquote(path))} "{_apache_target(r["new_normalized"])}" [R={r["redirect_status"]},L,NE,QSD]')
    if plain:
        lines.append("# Plain paths. A query string on the visitor's request is passed on to the new address.")
        for r in plain:
            lines.append(f'RewriteRule {pattern_for(unquote(r["old_path_key"]))} "{_apache_target(r["new_normalized"])}" [R={r["redirect_status"]},L,NE]')
    lines.append("</IfModule>")
    return ("\n".join(lines) + "\n").encode("utf-8")


def _nginx_string(value):
    """For a map VALUE: nginx expands $variables inside values, so a literal $ must be percent-encoded."""
    return value.replace("\\", "%5C").replace('"', "%22").replace("$", "%24").replace("\n", "").replace("\r", "")


def _nginx_key(value):
    """For a map KEY: nginx treats keys literally (verified with nginx 1.24), so a raw $ must stay as it is.
    Browsers send $ unencoded, and nginx compares against the raw request line."""
    return value.replace("\\", "%5C").replace('"', "%22").replace("\n", "").replace("\r", "")


def build_nginx(rows, project_name, strip_trailing_slash=True):
    """nginx: one `map` per redirect status code, mapping the exact request path to the destination.
    nginx cannot take a variable status code, so each code gets its own map and its own `if` line."""
    codes = sorted({r["redirect_status"] for r in rows})
    lines = [
        "# WebPuente redirect maps - REVIEW BEFORE USE (see the checklist in the project docs).",
        f"# Project: {_apache_comment(project_name)}",
        f"# Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}",
        "# 1) include this file inside the http { } block",
        "# 2) inside the server { } block add these lines:",
    ]
    for code in codes:
        lines.append(f"#      if ($webpuente_r{code}) {{ return {code} $webpuente_r{code}; }}")
    lines += ["# Large lists may need:  map_hash_bucket_size 128;  and  map_hash_max_size 8192;", ""]
    seen = set()
    for code in codes:
        lines += [f"map $request_uri $webpuente_r{code} {{", '    default "";']
        for r in rows:
            if r["redirect_status"] != code:
                continue
            variants = [r["old_path_key"]]
            if strip_trailing_slash and "?" not in r["old_path_key"] and r["old_path_key"] != "/":
                variants.append(r["old_path_key"] + "/")  # normalisation removed the slash, so match both forms
            expanded = []
            for variant in variants:
                expanded.append(variant)
                if "$" in variant:
                    expanded.append(variant.replace("$", "%24"))  # a client may also send it encoded
            for variant in expanded:
                if variant in seen:
                    continue
                seen.add(variant)
                lines.append(f'    "{_nginx_key(variant)}" "{_nginx_string(r["new_normalized"])}";')
        lines += ["}", ""]
    return ("\n".join(lines)).encode("utf-8")


# ---------------------------------------------------------------- light syntax checks (no server needed)
_AQ = r'"(?:[^"\\\n]|\\.)*"'  # a double-quoted Apache string; backslash escapes (\" \\ \. \%) are allowed inside
APACHE_LINE = re.compile(
    r"^(#.*|RewriteCond %\{QUERY_STRING\} " + _AQ
    + r"|RewriteRule (?:\\.|\S)+ " + _AQ + r" \[R=\d{3},L,NE(?:,QSD)?\]|<IfModule mod_rewrite\.c>|</IfModule>|RewriteEngine On|)$"
)
NGINX_LINE = re.compile(r'^(#.*|map \$request_uri \$webpuente_r\d{3} \{|    default "";|    "[^"\n]*" "[^"\n]*";|\}|)$')


def validate_syntax(fmt, content):
    """Returns a list of problems (empty = looks fine). This is a shape check, not a replacement for
    `apachectl configtest` / `nginx -t` on the real server."""
    text = content.decode("utf-8")
    pattern = APACHE_LINE if fmt == "htaccess" else NGINX_LINE
    problems = [f"line {i}: unexpected format" for i, line in enumerate(text.split("\n"), 1) if not pattern.match(line)]
    if fmt == "nginx" and text.count("{") != text.count("}"):
        problems.append("unbalanced braces")
    return problems
